from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from orders.models import Order


def generate_orders_excel(orders, title="Orders Report"):
    """Orders queryset ko ek formatted .xlsx Workbook me convert karta hai — accounting/tracking ke liye."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Orders"

    headers = [
        "Bill Number", "Invoice Number", "Customer", "Phone", "Email",
        "Status", "Subtotal", "CGST", "SGST", "Total Amount",
        "Payment Method", "Reward Points", "Order Date", "Delivered At", "Paid At",
    ]
    ws.append(headers)

    header_fill = PatternFill(start_color="1a1512", end_color="1a1512", fill_type="solid")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill

    total_amount_sum = 0
    for order in orders:
        invoice = getattr(order, "invoice", None)
        payment_detail = getattr(order, "payment_detail", None)
        ws.append([
            order.bill_number,
            invoice.invoice_number if invoice else "",
            order.user.get_full_name() or order.user.username,
            order.user.phone or "",
            order.user.email,
            order.get_status_display(),
            float(invoice.subtotal) if invoice else float(order.total_amount),
            float(invoice.cgst_amount) if invoice else 0,
            float(invoice.sgst_amount) if invoice else 0,
            float(invoice.total_amount) if invoice else float(order.total_amount),
            payment_detail.get_method_display() if payment_detail else "Not Paid",
            order.reward_points,
            order.created_at.strftime("%Y-%m-%d %H:%M"),
            order.delivered_at.strftime("%Y-%m-%d %H:%M") if order.delivered_at else "",
            order.paid_at.strftime("%Y-%m-%d %H:%M") if order.paid_at else "",
        ])
        total_amount_sum += float(invoice.total_amount) if invoice else float(order.total_amount)

    # Totals row
    ws.append([])
    total_row = ["", "", "", "", "", "TOTAL", "", "", "", total_amount_sum, "", "", "", "", ""]
    ws.append(total_row)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)

    for col in ws.columns:
        max_len = max((len(str(c.value)) for c in col if c.value is not None), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 3, 35)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
